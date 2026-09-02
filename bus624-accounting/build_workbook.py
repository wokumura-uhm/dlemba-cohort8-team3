#!/usr/bin/env python3
"""
BUS 624 - Comprehensive Financial Statement Analysis
Step 2 of 2: turn edgar_facts.csv into the Excel deliverable.

DESIGN PRINCIPLE
    Nothing downstream is a hardcoded number. Every common-size percentage and
    every ratio is a live Excel formula pointing back at the Source tab. That
    matters for two reasons:
      1. Antal wants you to verify by hand. When you check a figure against the
         10-K and it is wrong, you type the right number in ONE cell on Source
         and all 200-odd downstream cells re-compute. You never re-do the sheet.
      2. When she asks "where did this ratio come from?", you click the cell and
         the formula shows her. That is the whole ballgame on this assignment.

USAGE
    pip install openpyxl
    python3 pull_edgar.py --years 4          # 4, not 3 - see note below
    python3 build_workbook.py                # -> WMT_FSA_Workbook.xlsx

WHY FOUR YEARS AND NOT THREE
    The assignment wants 3 years of ratios. But ROA, ROE, and every turnover
    ratio use AVERAGE balances - (opening + closing) / 2. The oldest of your
    three years needs an opening balance, which is the year before it. So you
    need 4 years of balance sheet data to produce 3 years of ratios.
"""
import csv, os, sys
from collections import OrderedDict

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("Need openpyxl:  pip install openpyxl")

CSV_IN  = "edgar_facts.csv"
XLSX_OUT = "WMT_FSA_Workbook.xlsx"

SUBJECT = "WMT"
COMPS   = ["COST", "TGT", "KR"]
ALL_CO  = [SUBJECT] + COMPS
CO_NAME = {"WMT": "Walmart", "COST": "Costco", "TGT": "Target", "KR": "Kroger"}

# ---- styling --------------------------------------------------------------
H1    = Font(bold=True, size=14)
H2    = Font(bold=True, size=11, color="FFFFFF")
BOLD  = Font(bold=True)
ITAL  = Font(italic=True, size=9, color="666666")
FILL  = PatternFill("solid", fgColor="1F3864")
SUBF  = PatternFill("solid", fgColor="D9E2F3")
WARNF = PatternFill("solid", fgColor="FFF2CC")
THIN  = Border(bottom=Side(style="thin", color="BFBFBF"))
PCT   = '0.0%'
NUM   = '#,##0'
RAT   = '0.00'


def hdr(ws, row, labels, start=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=start + i, value=t)
        c.font, c.fill, c.alignment = H2, FILL, Alignment(horizontal="center", wrap_text=True)
    return row + 1


# ===========================================================================
# 1. Source tab - raw EDGAR facts + the hand-verification columns
# ===========================================================================
def build_source(wb, rows):
    ws = wb.create_sheet("Source")
    ws["A1"] = "Source data - pulled from SEC EDGAR XBRL (companyfacts API)"
    ws["A1"].font = H1
    ws["A2"] = ("Every row carries the accession number it came from. To verify: open "
                "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany and search the accession, "
                "or go to the company's 10-K directly.")
    ws["A2"].font = ITAL
    ws["A3"] = ("VERIFICATION PROTOCOL -> Check the EDGAR value against the printed 10-K. If it matches, "
                "put your initials in 'Verified by'. If it does NOT match, type the correct number in "
                "'Hand-verified value' and the whole workbook re-computes off your number.")
    ws["A3"].font = Font(italic=True, size=9, color="C00000")

    cols = ["ticker", "concept", "us_gaap_tag", "unit", "fiscal_year", "period_start",
            "period_end", "value", "form", "filed", "accession",
            "Hand-verified value", "Verified by", "Note / 10-K page", "Effective value", "key"]
    r = hdr(ws, 5, cols)

    first_data = r
    for row in rows:
        ws.cell(row=r, column=1,  value=row["ticker"])
        ws.cell(row=r, column=2,  value=row["concept"])
        ws.cell(row=r, column=3,  value=row["us_gaap_tag"])
        ws.cell(row=r, column=4,  value=row["unit"])
        ws.cell(row=r, column=5,  value=int(row["fiscal_year"]))
        ws.cell(row=r, column=6,  value=row["period_start"])
        ws.cell(row=r, column=7,  value=row["period_end"])
        v = ws.cell(row=r, column=8, value=float(row["value"])); v.number_format = NUM
        ws.cell(row=r, column=9,  value=row["form"])
        ws.cell(row=r, column=10, value=row["filed"])
        ws.cell(row=r, column=11, value=row["accession"])
        ws.cell(row=r, column=12).fill = WARNF          # hand-verified value (blank)
        ws.cell(row=r, column=13).fill = WARNF          # verified by
        ws.cell(row=r, column=14).fill = WARNF          # note
        e = ws.cell(row=r, column=15, value=f'=IF(L{r}<>"",L{r},H{r})'); e.number_format = NUM
        e.font = BOLD
        ws.cell(row=r, column=16, value=f'=A{r}&"|"&B{r}&"|"&E{r}')
        r += 1
    last_data = r - 1

    widths = [8, 26, 46, 10, 7, 12, 12, 18, 7, 12, 22, 18, 12, 28, 18, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A6"
    return first_data, last_data


# ===========================================================================
# 2. Data grid - concept x company x year, via INDEX/MATCH on the key column
# ===========================================================================
CONCEPT_ORDER = [
    ("INCOME STATEMENT", None),
    ("Revenue", "Revenue"), ("COGS", "COGS"), ("GrossProfit", "Gross profit"),
    ("SGA", "SG&A"), ("OperatingIncome", "Operating income"),
    ("InterestExpense", "Interest expense"), ("PretaxIncome", "Pre-tax income"),
    ("TaxExpense", "Income tax expense"), ("NetIncome", "Net income"),
    ("EPSDiluted", "EPS - diluted"), ("SharesDiluted", "Diluted shares"),
    ("BALANCE SHEET", None),
    ("Cash", "Cash & equivalents"), ("Receivables", "Receivables, net"),
    ("Inventory", "Inventory"), ("CurrentAssets", "Total current assets"),
    ("PPENet", "PP&E, net"), ("Goodwill", "Goodwill"), ("TotalAssets", "TOTAL ASSETS"),
    ("AccountsPayable", "Accounts payable"), ("CurrentLiabilities", "Total current liabilities"),
    ("LongTermDebtCurrent", "LT debt - current portion"), ("LongTermDebt", "Long-term debt"),
    ("OperatingLeaseLiabCurrent", "Operating lease liab - current"),
    ("OperatingLeaseLiabNoncurrent", "Operating lease liab - noncurrent"),
    ("TotalLiabilities", "TOTAL LIABILITIES"), ("Equity", "Stockholders' equity"),
    ("CASH FLOW", None),
    ("CFO", "Cash from operations"), ("CFI", "Cash from investing"),
    ("CFF", "Cash from financing"), ("CapEx", "Capital expenditures"),
    ("DepAmort", "Depreciation & amortization"), ("DividendsPaid", "Dividends paid"),
    ("Buybacks", "Share repurchases"),
]


def build_grid(wb, years, src_first, src_last):
    ws = wb.create_sheet("Data")
    ws["A1"] = "Consolidated data grid - all formulas, sourced from the Source tab"
    ws["A1"].font = H1
    ws["A2"] = "Do not type numbers here. Fix errors on the Source tab instead."
    ws["A2"].font = ITAL

    rng = f"Source!$P${src_first}:$P${src_last}"
    val = f"Source!$O${src_first}:$O${src_last}"

    # header: company blocks, each with its fiscal years
    ws.cell(row=4, column=1, value="Line item").font = H2
    ws.cell(row=4, column=1).fill = FILL
    col = 2
    colmap = {}
    for co in ALL_CO:
        c = ws.cell(row=4, column=col, value=CO_NAME[co])
        c.font, c.fill, c.alignment = H2, FILL, Alignment(horizontal="center")
        ws.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + len(years) - 1)
        for j, y in enumerate(years):
            yc = ws.cell(row=5, column=col + j, value=f"FY{y}")
            yc.font, yc.fill, yc.alignment = BOLD, SUBF, Alignment(horizontal="center")
            colmap[(co, y)] = col + j
        col += len(years)

    r = 6
    rowmap = {}
    for concept, label in CONCEPT_ORDER:
        if label is None:
            c = ws.cell(row=r, column=1, value=concept)
            c.font, c.fill = BOLD, SUBF
            for cc in range(2, col):
                ws.cell(row=r, column=cc).fill = SUBF
            r += 1
            continue
        ws.cell(row=r, column=1, value=label)
        rowmap[concept] = r
        for co in ALL_CO:
            for y in years:
                key = f'"{co}|{concept}|{y}"'
                f = f'=IFERROR(INDEX({val},MATCH({key},{rng},0)),"")'
                cell = ws.cell(row=r, column=colmap[(co, y)], value=f)
                cell.number_format = RAT if concept.startswith("EPS") else NUM
        r += 1

    ws.column_dimensions["A"].width = 34
    for cc in range(2, col):
        ws.column_dimensions[get_column_letter(cc)].width = 15
    ws.freeze_panes = "B6"
    return rowmap, colmap, years


# ===========================================================================
# 3. Common-size tabs - vertical (% of base) and horizontal (YoY growth)
# ===========================================================================
IS_ITEMS = ["Revenue", "COGS", "GrossProfit", "SGA", "OperatingIncome",
            "InterestExpense", "PretaxIncome", "TaxExpense", "NetIncome"]
BS_ITEMS = ["Cash", "Receivables", "Inventory", "CurrentAssets", "PPENet", "Goodwill",
            "TotalAssets", "AccountsPayable", "CurrentLiabilities", "LongTermDebtCurrent",
            "LongTermDebt", "OperatingLeaseLiabCurrent", "OperatingLeaseLiabNoncurrent",
            "TotalLiabilities", "Equity"]


def build_common_size(wb, title, items, base_concept, rowmap, colmap, years, note):
    ws = wb.create_sheet(title)
    ws["A1"] = f"{title} - combination vertical and horizontal analysis"
    ws["A1"].font = H1
    ws["A2"] = note
    ws["A2"].font = ITAL

    label_of = {c: l for c, l in CONCEPT_ORDER if l}

    r = 4
    for co in ALL_CO:
        ws.cell(row=r, column=1, value=CO_NAME[co]).font = H1
        r += 1
        heads = ["Line item"] + [f"FY{y} % of base" for y in years] \
                              + [f"FY{y} YoY growth" for y in years[1:]]
        r = hdr(ws, r, heads)
        base_row = rowmap[base_concept]
        for it in items:
            ws.cell(row=r, column=1, value=label_of.get(it, it))
            src = rowmap[it]
            for j, y in enumerate(years):
                dc = get_column_letter(colmap[(co, y)])
                f = f'=IFERROR(Data!{dc}{src}/Data!{dc}{base_row},"")'
                c = ws.cell(row=r, column=2 + j, value=f); c.number_format = PCT
            for j, y in enumerate(years[1:]):
                cur = get_column_letter(colmap[(co, y)])
                pri = get_column_letter(colmap[(co, years[years.index(y) - 1])])
                f = f'=IFERROR(Data!{cur}{src}/Data!{pri}{src}-1,"")'
                c = ws.cell(row=r, column=2 + len(years) + j, value=f); c.number_format = PCT
            r += 1
        r += 2

    ws.column_dimensions["A"].width = 34
    for cc in range(2, 2 + len(years) * 2):
        ws.column_dimensions[get_column_letter(cc)].width = 16
    return ws


# ===========================================================================
# 4. Ratios tab
# ===========================================================================
# (label, category, formula builder, number format)
# cur(concept) -> current year cell ; avg(concept) -> average balance cell
RATIOS = [
    ("Gross profit margin",        "Profitability", lambda c, a: f'{c("GrossProfit")}/{c("Revenue")}', PCT),
    ("Operating profit margin",    "Profitability", lambda c, a: f'{c("OperatingIncome")}/{c("Revenue")}', PCT),
    ("Net profit margin",          "Profitability", lambda c, a: f'{c("NetIncome")}/{c("Revenue")}', PCT),
    ("Return on assets (ROA)",     "Profitability", lambda c, a: f'{c("NetIncome")}/{a("TotalAssets")}', PCT),
    ("Return on equity (ROE)",     "Profitability", lambda c, a: f'{c("NetIncome")}/{a("Equity")}', PCT),
    ("Effective tax rate",         "Profitability", lambda c, a: f'{c("TaxExpense")}/{c("PretaxIncome")}', PCT),
    ("Current ratio",              "Liquidity",     lambda c, a: f'{c("CurrentAssets")}/{c("CurrentLiabilities")}', RAT),
    ("Quick ratio",                "Liquidity",     lambda c, a: f'({c("Cash")}+{c("Receivables")})/{c("CurrentLiabilities")}', RAT),
    ("Operating cash flow to current liabilities", "Liquidity", lambda c, a: f'{c("CFO")}/{a("CurrentLiabilities")}', RAT),
    ("Asset turnover",             "Efficiency",    lambda c, a: f'{c("Revenue")}/{a("TotalAssets")}', RAT),
    ("Inventory turnover",         "Efficiency",    lambda c, a: f'{c("COGS")}/{a("Inventory")}', RAT),
    ("Days inventory outstanding", "Efficiency",    lambda c, a: f'365/({c("COGS")}/{a("Inventory")})', RAT),
    ("Days payable outstanding",   "Efficiency",    lambda c, a: f'365/({c("COGS")}/{a("AccountsPayable")})', RAT),
    ("Fixed asset turnover",       "Efficiency",    lambda c, a: f'{c("Revenue")}/{a("PPENet")}', RAT),
    ("Total liabilities to equity","Solvency",      lambda c, a: f'{c("TotalLiabilities")}/{c("Equity")}', RAT),
    ("Total debt to equity",       "Solvency",      lambda c, a: f'({c("LongTermDebt")}+{c("LongTermDebtCurrent")})/{c("Equity")}', RAT),
    ("Times interest earned",      "Solvency",      lambda c, a: f'{c("OperatingIncome")}/{c("InterestExpense")}', RAT),
    ("Free cash flow ($M)",        "Cash",          lambda c, a: f'{c("CFO")}-{c("CapEx")}', NUM),
    ("Capex to revenue",           "Cash",          lambda c, a: f'{c("CapEx")}/{c("Revenue")}', PCT),
    ("Dividend payout ratio",      "Cash",          lambda c, a: f'{c("DividendsPaid")}/{c("NetIncome")}', PCT),
]

DUPONT = [
    ("Net profit margin",     lambda c, a: f'{c("NetIncome")}/{c("Revenue")}', PCT),
    ("x Asset turnover",      lambda c, a: f'{c("Revenue")}/{a("TotalAssets")}', RAT),
    ("= Return on assets",    lambda c, a: f'{c("NetIncome")}/{a("TotalAssets")}', PCT),
    ("x Financial leverage",  lambda c, a: f'{a("TotalAssets")}/{a("Equity")}', RAT),
    ("= Return on equity",    lambda c, a: f'{c("NetIncome")}/{a("Equity")}', PCT),
]


def build_ratios(wb, rowmap, colmap, years, all_years):
    ws = wb.create_sheet("Ratios")
    ws["A1"] = "Ratio analysis"
    ws["A1"].font = H1
    ws["A2"] = ("Assignment asks for 15 ratios that BEST SUPPORT YOUR STRATEGY. There are 20 here on "
                "purpose - that is a menu, not an answer. Delete the 5 that do not carry your argument. "
                "Ratios using average balances need an opening balance, so the earliest year shown is "
                "one year later than the earliest year of raw data.")
    ws["A2"].font = ITAL
    ws["A3"] = "Industry average column is MANUAL - pull it from Mergent Online via the UH Hamilton Library and type it in."
    ws["A3"].font = Font(italic=True, size=9, color="C00000")

    def make_ref(co, y, avg=False):
        def cur(concept):
            return f'Data!{get_column_letter(colmap[(co, y)])}{rowmap[concept]}'
        def average(concept):
            pi = all_years.index(y) - 1
            if pi < 0:
                return cur(concept)
            prev = f'Data!{get_column_letter(colmap[(co, all_years[pi])])}{rowmap[concept]}'
            return f'AVERAGE({cur(concept)},{prev})'
        return cur, average

    r = 5
    heads = ["Ratio", "Category"]
    for co in ALL_CO:
        heads += [f"{CO_NAME[co]} FY{y}" for y in years]
    heads += ["Industry avg (manual)", "Source of industry avg"]
    r = hdr(ws, r, heads)

    ind_col = 3 + len(ALL_CO) * len(years)

    def emit(rows, r):
        for item in rows:
            label, fn, fmt = (item[0], item[-2], item[-1])
            cat = item[1] if len(item) == 4 else ""
            ws.cell(row=r, column=1, value=label).border = THIN
            ws.cell(row=r, column=2, value=cat)
            cc = 3
            for co in ALL_CO:
                for y in years:
                    cur, avg = make_ref(co, y)
                    c = ws.cell(row=r, column=cc, value=f'=IFERROR({fn(cur, avg)},"")')
                    c.number_format = fmt
                    cc += 1
            ws.cell(row=r, column=ind_col).fill = WARNF
            ws.cell(row=r, column=ind_col + 1).fill = WARNF
            r += 1
        return r

    r = emit(RATIOS, r)

    r += 2
    ws.cell(row=r, column=1, value="DuPont decomposition of ROE  (syllabus requires this explicitly)").font = H1
    r += 1
    r = hdr(ws, r, heads)
    r = emit([(l, "DuPont", f, fmt) for l, f, fmt in DUPONT], r)

    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 14
    for cc in range(3, ind_col + 2):
        ws.column_dimensions[get_column_letter(cc)].width = 15
    ws.freeze_panes = "C6"
    return ws


# ===========================================================================
# 5. README + takeaways scaffold
# ===========================================================================
def build_readme(wb, years):
    ws = wb.create_sheet("READ ME FIRST", 0)
    ws.column_dimensions["A"].width = 112
    lines = [
        ("BUS 624 - Walmart Comprehensive Financial Statement Analysis", H1),
        ("", None),
        ("WHAT THIS WORKBOOK IS", BOLD),
        ("Raw XBRL facts pulled from SEC EDGAR for Walmart, Costco, Target and Kroger, plus", None),
        ("common-size and ratio analysis computed as live formulas off those facts.", None),
        ("", None),
        ("HOW THE TABS RELATE", BOLD),
        ("  Source    Raw EDGAR facts. The ONLY tab where you type numbers.", None),
        ("  Data      Formula grid. Pulls from Source by ticker|concept|year. Never type here.", None),
        ("  CS-Income Common-size income statement, % of revenue + YoY growth.", None),
        ("  CS-Balance Common-size balance sheet, % of total assets + YoY growth.", None),
        ("  Ratios    20 candidate ratios + DuPont. Pick the 15 that carry your argument.", None),
        ("  Takeaways Where you write the 2-3 findings each presentation requires.", None),
        ("", None),
        ("THE VERIFICATION RULE - READ THIS", Font(bold=True, color="C00000")),
        ("Antal's instruction: AI output is a tool 'much like an analyst report, rather than an", None),
        ("end product', and you must pull the original source documents and confirm accuracy.", None),
        ("So: this workbook is UNVERIFIED until a human checks it. Every Source row carries the", None),
        ("accession number of the filing it came from. Open that filing, find the number, and", None),
        ("initial the 'Verified by' column. If EDGAR disagrees with the printed 10-K, type the", None),
        ("10-K number into 'Hand-verified value' - the workbook then computes off your number.", None),
        ("", None),
        ("KNOWN TRAPS", BOLD),
        ("1. Costco's fiscal year ends in late August/early September; Walmart, Target and Kroger", None),
        ("   end in late January/early February. So Costco 'FY2026' covers a different calendar", None),
        ("   window than Walmart 'FY2026'. Say so out loud in your presentation - it is exactly", None),
        ("   the kind of comparability caveat that earns marks.", None),
        ("2. Kroger and Costco both run fuel businesses; Target does not. That asymmetry matters", None),
        ("   for margin comparison, because fuel is high-revenue and low-margin, so it structurally", None),
        ("   depresses gross margin percentages without saying anything about merchandising skill.", None),
        ("3. Operating lease liabilities are large for all four and sit outside 'long-term debt'.", None),
        ("   A debt-to-equity ratio that ignores them understates leverage for retailers.", None),
        ("4. Blank cells are a FINDING, not a failure. They usually mean the company reports that", None),
        ("   item on a differently-named line. Go look, then record what you found.", None),
        ("", None),
        ("REBUILDING FROM SCRATCH", BOLD),
        ("  python3 pull_edgar.py --years 4     # 4 years of data -> 3 years of average-balance ratios", None),
        ("  python3 build_workbook.py", None),
        ("Re-running overwrites the workbook, including your hand-verification notes. Save a copy", None),
        ("under a new name once you start verifying.", None),
    ]
    for i, (t, f) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=t)
        if f:
            c.font = f
    return ws


def build_takeaways(wb):
    ws = wb.create_sheet("Takeaways")
    ws["A1"] = "Presentation take-aways - the assignment asks for 2-3 per section"
    ws["A1"].font = H1
    ws["A2"] = ("Antal is explicit: 'I am looking for what you add to the picture, not what I can get "
                "from AI.' The workbook produces the numbers. This tab is where the numbers become an "
                "argument. Nobody can write this for you.")
    ws["A2"].font = ITAL
    r = hdr(ws, 4, ["Deliverable", "#", "Take-away (one sentence, must be falsifiable)",
                    "Which number supports it (tab + cell)", "So what? (why the Budget Committee cares)", "Owner"])
    blocks = [("Part 1 - MD&A / Message from Co / Auditor reports", 3),
              ("Part 2 - Common-size analysis", 3),
              ("Part 3 - Ratio analysis", 3),
              ("Part 4 - Strategy & recommendations", 3),
              ("Part 4 - Negative impacts & mitigation", 3)]
    for name, n in blocks:
        for i in range(1, n + 1):
            ws.cell(row=r, column=1, value=name if i == 1 else "")
            ws.cell(row=r, column=2, value=i)
            for cc in range(3, 7):
                ws.cell(row=r, column=cc).fill = WARNF
            r += 1
        r += 1
    for cc, w in zip(range(1, 7), [42, 4, 56, 30, 48, 12]):
        ws.column_dimensions[get_column_letter(cc)].width = w
    return ws


# ===========================================================================
def main():
    if not os.path.exists(CSV_IN):
        sys.exit(f"Missing {CSV_IN}. Run:  python3 pull_edgar.py --years 4")

    with open(CSV_IN) as f:
        rows = list(csv.DictReader(f))
    if not rows:
        sys.exit(f"{CSV_IN} is empty.")

    all_years = sorted({int(r["fiscal_year"]) for r in rows})
    ratio_years = all_years[1:] if len(all_years) > 3 else all_years
    ratio_years = ratio_years[-3:]

    wb = Workbook()
    wb.remove(wb.active)

    build_readme(wb, all_years)
    sf, sl = build_source(wb, rows)
    rowmap, colmap, _ = build_grid(wb, all_years, sf, sl)
    build_common_size(wb, "CS-Income", IS_ITEMS, "Revenue", rowmap, colmap, all_years,
                      "Vertical analysis is % of revenue. Horizontal analysis is year-over-year growth.")
    build_common_size(wb, "CS-Balance", BS_ITEMS, "TotalAssets", rowmap, colmap, all_years,
                      "Vertical analysis is % of total assets. Horizontal analysis is year-over-year growth.")
    build_ratios(wb, rowmap, colmap, ratio_years, all_years)
    build_takeaways(wb)

    wb.save(XLSX_OUT)
    print(f"Wrote {XLSX_OUT}")
    print(f"  companies : {', '.join(CO_NAME[c] for c in ALL_CO)}")
    print(f"  data years: {all_years}")
    print(f"  ratio yrs : {ratio_years}  (earliest dropped - needed as opening balance)")
    print("\nNext: open it, go to the Source tab, and start verifying against the 10-Ks.")


if __name__ == "__main__":
    main()
