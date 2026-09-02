#!/usr/bin/env python3
"""
Generates the deck shells and the memo skeleton for the BUS 624 group project.

THESE ARE STRUCTURE ONLY - deliberately. Every slide and every heading carries a
prompt telling you what belongs there and what Antal is grading, but no content.
The analysis has to be the group's own: "I am looking for what you add to the
picture, not what I can get from AI."
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from docx import Document
from docx.shared import Pt as DPt, RGBColor as DRGB
from docx.enum.text import WD_ALIGN_PARAGRAPH

NAVY = RGBColor(0x1F, 0x38, 0x64)

# (slide title, [bullet prompts], speaker-note prompt)
DECKS = {
 "Part1_Overview_MDA_Auditor": [
  ("Walmart Inc. (NYSE: WMT)\nPart 1 - Company, Message from the Company, MD&A, Auditor Reports",
   ["Richie · Halia · Sean · Wade · Micah", "Presented 13 October 2026 · Lead: Halia", "BUS 624 - Fall 2026"],
   "Antal's framing: present as if answering your next-door neighbour. Do not read from the 10-K."),
  ("What Walmart does",
   ["Business model in plain language", "Segments and what each contributes",
    "Two or three genuinely interesting 'to-knows'"],
   "Plain language. If a slide needs a finance degree to follow, rewrite it."),
  ("Our working thesis",
   ["The uncertainty we are preparing the Budget Committee for:",
    "Our preliminary strategic direction:",
    "Why this matters to Walmart specifically"],
   "CRITICAL: the syllabus wants Part 1 takeaways chosen because they SUPPORT your "
   "strategy. So the strategy has to exist before this presentation. State it as a "
   "hypothesis you may revise - the syllabus explicitly permits that."),
  ("Message from the Company - takeaways",
   ["Takeaway 1 - and which recommendation it supports",
    "Takeaway 2 - and which recommendation it supports",
    "Takeaway 3 (optional)"],
   "Two to three only. Each must connect to the strategy, not just be interesting."),
  ("MD&A - takeaways",
   ["Takeaway 1 - and which recommendation it supports",
    "Takeaway 2 - and which recommendation it supports",
    "Takeaway 3 (optional)"],
   "MD&A is where management explains results in their own words. Quote sparingly, "
   "paraphrase mostly, and cite the page."),
  ("Auditor reports - takeaways",
   ["Internal control over financial reporting - opinion and any material weakness",
    "Financial statements - opinion type",
    "Critical Audit Matters: what the auditor flagged as hardest to judge"],
   "There are TWO reports - internal controls and financial statements. Antal says so "
   "explicitly. Critical Audit Matters are the highest-value part: they tell you where "
   "the accounting is most subjective."),
  ("So what for the Budget Committee",
   ["What Part 1 establishes", "What we test next in the common-size analysis"],
   "Close by pointing forward. It makes the four parts read as one project."),
 ],
 "Part2_CommonSize": [
  ("Part 2 - Common-Size Analysis\nWalmart vs Costco, Target, Kroger",
   ["Richie · Halia · Sean · Wade · Micah", "Presented 27 October 2026 · Lead: Richie", "BUS 624 - Fall 2026"], "Submitted Fri 23 Oct - two days ahead of the Sunday deadline, to leave room for Antal's feedback before presenting."),
  ("Method and comparability caveats",
   ["Vertical: income statement as % of revenue; balance sheet as % of total assets",
    "Horizontal: year-over-year growth, 3 years",
    "CAVEAT - Costco's fiscal year ends in August; the others end in Jan/Feb",
    "CAVEAT - Kroger and Costco sell fuel; Target does not"],
   "Leading with comparability limits reads as rigour, not weakness. Most groups skip this."),
  ("Common-size income statement",
   ["Chart or table", "Where Walmart's cost structure differs from peers"],
   "Do not narrate the whole table. Point at the two or three lines that matter."),
  ("Common-size balance sheet",
   ["Chart or table", "Asset intensity and how the balance sheet is funded"],
   "Watch operating lease liabilities - large for all four, and outside 'long-term debt'."),
  ("Trends: significant increases and decreases",
   ["What moved materially over 3 years", "What we think is driving it"],
   "'Significant' needs a threshold. Say what yours is."),
  ("Takeaways",
   ["Takeaway 1 - and the recommendation it supports",
    "Takeaway 2 - and the recommendation it supports",
    "Takeaway 3 (optional)"],
   "Two to three. Each tied to the strategy."),
  ("Strengths and weaknesses",
   ["Strengths the numbers actually show", "Weaknesses the numbers actually show"],
   "Every claim needs a cell reference behind it."),
 ],
 "Part3_Ratios": [
  ("Part 3 - Ratio Analysis\nWalmart vs Costco, Target, Kroger, and industry averages",
   ["Richie · Halia · Sean · Wade · Micah", "Presented 3 November 2026 · Lead: Sean", "BUS 624 - Fall 2026"], "Submitted Fri 30 Oct, ahead of the Sunday deadline."),
  ("The 15 ratios and why these 15",
   ["Profitability:", "Liquidity:", "Efficiency / turnover:", "Solvency:"],
   "The assignment says ratios that BEST SUPPORT YOUR STRATEGY. Be ready to justify why "
   "you dropped the ones you dropped - that is the actual test here."),
  ("DuPont decomposition of ROE",
   ["Net profit margin x asset turnover = ROA", "ROA x financial leverage = ROE",
    "Margins or volume? Capital intensive? Highly levered?"],
   "The syllabus names this explicitly. Answer the margin-vs-volume question out loud."),
  ("Walmart vs competitors",
   ["Where Walmart leads", "Where Walmart lags", "Where the gap is widening or closing"],
   "Direction of travel over 3 years matters more than any single year's level."),
  ("Walmart vs industry average",
   ["Source of industry data (Mergent Online, UH Hamilton Library)",
    "Where Walmart sits against the industry"],
   "Cite the industry source explicitly - she asks for it."),
  ("Takeaways",
   ["Takeaway 1 - and the recommendation it supports",
    "Takeaway 2 - and the recommendation it supports",
    "Takeaway 3 (optional)"],
   "Two to three. Tied to strategy."),
  ("Strengths, weaknesses, and adaptation",
   ["Strengths", "Weaknesses", "How Walmart is adapting to the changing environment"],
   "The adaptation question is asked directly in the syllabus. Do not skip it."),
 ],
 "Part4_Strategy_and_Impacts": [
  ("Part 4 - Strategies, Recommendations, and Potential Negative Impacts\nMemo to the Budget Committee",
   ["Richie · Halia · Sean · Wade · Micah", "Presented 10 November 2026 · Lead: Wade", "BUS 624 - Fall 2026"], "Draft submitted Fri 6 Nov, ahead of the Sunday deadline."),
  ("As-is: where Walmart stands today",
   ["Current position on the uncertainty we were asked about",
    "The financial evidence, from Parts 2 and 3"],
   "'As-is' must be evidenced from your own analysis, not from general knowledge."),
  ("What-if: the scenario",
   ["The scenario and its magnitude", "What it does to the financials",
    "Which line items move, and by how much"],
   "Quantify. A what-if without numbers is an opinion."),
  ("Strategic option 1",
   ["What we would do", "Initial investment $", "Training cost $", "Annual maintenance $",
    "Implementation timeline"],
   "The assignment requires AT LEAST TWO options compared. This is option 1."),
  ("Strategic option 2",
   ["What we would do", "Initial investment $", "Training cost $", "Annual maintenance $",
    "Implementation timeline"],
   "Option 2 must be a genuine alternative, not a strawman."),
  ("Recommendation",
   ["Which option and why", "SMART: specific, measurable, achievable, relevant, time-bound",
    "Total amount to invest", "Quantitative support", "Qualitative support"],
   "Check each SMART letter explicitly. Groups lose marks on 'measurable' most often."),
  ("Implementation timeline",
   ["Phase 1:", "Phase 2:", "Phase 3:", "Milestones and decision gates"],
   "A timeline with decision gates shows you thought about being wrong."),
  ("Potential negative impacts",
   ["Financial - e.g. operating cash reduced by $Y for N years",
    "Operational - what gets delayed or displaced",
    "Ethical - who is affected, including associates",
    "Why we selected these impacts"],
   "Antal gives the exact shape she wants: 'operating cash will be reduced by "
   "approximately $Y for 5 years, we will need to delay investment in Z for 2 years.' "
   "Match that specificity. The ethical dimension is graded - do not treat it as filler."),
  ("Mitigation",
   ["Mitigation for each impact above", "Our logic for why these mitigations work",
    "e.g. reskilling, phased investment, contingency planning"],
   "She names those three examples in the syllabus. Use them as a starting point, not a script."),
 ],
}


def build_deck(name, slides):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for i, (title, bullets, note) in enumerate(slides):
        layout = prs.slide_layouts[0] if i == 0 else prs.slide_layouts[1]
        s = prs.slides.add_slide(layout)
        s.shapes.title.text = title
        for p in s.shapes.title.text_frame.paragraphs:
            for r in p.runs:
                r.font.color.rgb = NAVY
                r.font.size = Pt(30 if i == 0 else 26)
                r.font.bold = True
        if i == 0:
            if len(s.placeholders) > 1:
                s.placeholders[1].text = "\n".join(bullets)
        else:
            tf = s.placeholders[1].text_frame
            tf.text = bullets[0]
            for b in bullets[1:]:
                tf.add_paragraph().text = b
            for p in tf.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(18)
        s.notes_slide.notes_text_frame.text = "PROMPT (delete before presenting): " + note
    out = f"{name}.pptx"
    prs.save(out)
    return out, len(slides)


MEMO = [
 ("MEMORANDUM", "h0", None),
 ("TO: Budget Committee, Walmart Inc.\nFROM: Richie, Halia, Sean, Wade, Micah Kosasa\nDATE: December 1, 2026\n"
  "RE: [Subject line naming the uncertainty and your recommendation]", "meta", None),
 ("Executive Summary", "h1",
  "Written LAST, read FIRST. One page maximum. State the ask, the finding, the recommendation, "
  "the dollar amount, and the main risk. A committee member who reads only this page should be "
  "able to vote."),
 ("Purpose", "h1", "Why this memo exists and what decision it is asking the Budget Committee to make."),
 ("Background", "h1",
  "Walmart's position, the uncertainty you were asked to prepare for, and the competitive context. "
  "Pull forward from Part 1 - and update it so it matches your FINAL recommendations, which the "
  "syllabus explicitly requires."),
 ("Discussion / Analysis", "h1",
  "The evidence. Common-size findings, ratio findings, competitor and industry comparison. "
  "Every number traced to a source. This is where Parts 2 and 3 earn their keep."),
 ("As-is: current state", "h2", "What are we doing now?"),
 ("What-if: scenario analysis", "h2", "What can we do to improve? Quantify the scenario."),
 ("Comparison of strategic options", "h2", "At least two genuine options, side by side, with costs."),
 ("Recommendations", "h1",
  "The syllabus is specific: explain the ADVANTAGES and DISADVANTAGES and how you plan to "
  "MINIMISE the disadvantages. Include the investment amount, the SMART framing, and the "
  "implementation timeline."),
 ("Recommended action and investment", "h2", "What, how much, by when."),
 ("Advantages", "h2", None),
 ("Disadvantages and downstream negative impacts", "h2",
  "Financial, operational, and ethical. Match Antal's specificity: '$Y for N years'."),
 ("Mitigation plan", "h2", "How each disadvantage above gets managed down."),
 ("Conclusion", "h1", "Short. Restate the decision you are asking for."),
 ("Appendix A - Individual Reflections", "h1",
  "ONE PER PERSON, WRITTEN INDIVIDUALLY. The syllabus requires this and it is easy to forget "
  "until the night before. Due on the 20 Nov call - it is on the tracker."),
 ("Appendix B - Sources", "h1",
  "Company website link and links to the 10-Ks - the syllabus requires these in the submission "
  "email to rantal@hawaii.edu."),
]


def build_memo():
    d = Document()
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", DPt(11)
    for text, kind, prompt in MEMO:
        if kind == "h0":
            p = d.add_paragraph(); r = p.add_run(text)
            r.bold, r.font.size = True, DPt(20); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif kind == "meta":
            for line in text.split("\n"):
                p = d.add_paragraph(); r = p.add_run(line); r.bold = True
            d.add_paragraph("_" * 78)
        elif kind == "h1":
            d.add_heading(text, level=1)
        else:
            d.add_heading(text, level=2)
        if prompt:
            p = d.add_paragraph()
            r = p.add_run("[ " + prompt + " ]")
            r.italic = True; r.font.size = DPt(9); r.font.color.rgb = DRGB(0x88, 0x88, 0x88)
            d.add_paragraph()
    d.save("Final_Memo_Skeleton.docx")
    return "Final_Memo_Skeleton.docx"


if __name__ == "__main__":
    for name, slides in DECKS.items():
        f, n = build_deck(name, slides)
        print(f"  {f}  ({n} slides, speaker-note prompts on every slide)")
    print(f"  {build_memo()}")
